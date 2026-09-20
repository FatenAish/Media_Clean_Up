import base64
import io
import os
import re
import time
import tempfile
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit

import imagehash
import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageOps
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

st.set_page_config(page_title="Shutterstock Article Image Checker", page_icon="🔎", layout="wide")

BASE_DIR = Path(__file__).resolve().parent
PROFILE_DIR = BASE_DIR / "work" / "shutterstock_chrome_profile_v2"
PROFILE_DIR.mkdir(parents=True, exist_ok=True)

MAX_RESULTS = 40
MATCH_THRESHOLD = 0.86
POSSIBLE_THRESHOLD = 0.76


def clean_url(url):
    url = str(url or "").strip()
    if not url:
        return ""
    p = urlsplit(url)
    return urlunsplit((p.scheme, p.netloc, p.path, "", ""))


def pil_from_bytes(data):
    return Image.open(io.BytesIO(data)).convert("RGB")


def normalize_image(data):
    img = ImageOps.exif_transpose(pil_from_bytes(data))
    img.thumbnail((3900, 3900), Image.Resampling.LANCZOS)
    for quality in (92, 88, 84, 80, 76, 72, 68):
        out = io.BytesIO()
        img.save(out, "JPEG", quality=quality, optimize=True)
        value = out.getvalue()
        if len(value) < 4_800_000:
            return value
    img.thumbnail((3000, 3000), Image.Resampling.LANCZOS)
    out = io.BytesIO()
    img.save(out, "JPEG", quality=65, optimize=True)
    return out.getvalue()


def crop_variant(img, ratio=0.84):
    w, h = img.size
    nw, nh = int(w * ratio), int(h * ratio)
    left = max(0, (w - nw) // 2)
    top = max(0, (h - nh) // 2)
    return img.crop((left, top, left + nw, top + nh))


def similarity(a_bytes, b_bytes):
    try:
        a = ImageOps.exif_transpose(pil_from_bytes(a_bytes))
        b = ImageOps.exif_transpose(pil_from_bytes(b_bytes))
        ac = crop_variant(a)
        bc = crop_variant(b)
        pairs = [
            (imagehash.phash(a, hash_size=16), imagehash.phash(b, hash_size=16)),
            (imagehash.dhash(a, hash_size=16), imagehash.dhash(b, hash_size=16)),
            (imagehash.whash(a, hash_size=16), imagehash.whash(b, hash_size=16)),
            (imagehash.phash(ac, hash_size=16), imagehash.phash(bc, hash_size=16)),
            (imagehash.dhash(ac, hash_size=16), imagehash.dhash(bc, hash_size=16)),
        ]
        scores = [max(0.0, 1.0 - ((x - y) / 256.0)) for x, y in pairs]
        scores.sort(reverse=True)
        return min(1.0, 0.55 * scores[0] + 0.25 * scores[1] + 0.20 * (sum(scores) / len(scores)))
    except Exception:
        return 0.0


def extract_article_images(page, article_url):
    page.goto(article_url, wait_until="domcontentloaded", timeout=90000)
    page.wait_for_timeout(2500)
    title = page.title()

    for _ in range(8):
        page.mouse.wheel(0, 1100)
        page.wait_for_timeout(250)

    candidates = []
    og = page.locator('meta[property="og:image"]')
    if og.count():
        src = og.get_attribute("content")
        if src:
            candidates.append({"src": urljoin(article_url, src), "width": 1200, "height": 600, "position": "Feature"})

    selectors = ["article img", "main article img", ".entry-content img", ".post-content img", ".article-content img"]
    selector = next((s for s in selectors if page.locator(s).count()), "main img")
    raw = page.locator(selector).evaluate_all(
        """imgs => imgs.map((img, idx) => ({
            src: img.currentSrc || img.src || img.getAttribute('data-src') || img.getAttribute('data-lazy-src') || img.getAttribute('data-original') || '',
            width: img.naturalWidth || img.width || 0,
            height: img.naturalHeight || img.height || 0,
            idx: idx + 1
        }))"""
    )

    for item in raw:
        src = str(item.get("src") or "").strip()
        if src and not src.startswith("data:"):
            candidates.append({
                "src": urljoin(article_url, src),
                "width": int(item.get("width") or 0),
                "height": int(item.get("height") or 0),
                "position": f"Body {item.get('idx', '')}",
            })

    excluded = (
        "logo", "icon", "avatar", "emoji", "favicon", "sprite", "placeholder",
        "author", "profile", "social", "facebook", "instagram", "linkedin",
        "youtube", "twitter", "whatsapp", "download-bayut-app", "app-store", "google-play",
    )

    seen = set()
    final = []
    for item in candidates:
        src = clean_url(item["src"])
        if not src or src in seen:
            continue
        if any(term in src.lower() for term in excluded):
            continue
        if item["width"] and item["width"] < 450:
            continue
        if item["height"] and item["height"] < 220:
            continue
        seen.add(src)
        item["src"] = src
        final.append(item)

    return title, final


def download_image(url):
    r = requests.get(url, headers={"User-Agent": "Mozilla/5.0", "Referer": "https://www.bayut.com/"}, timeout=45)
    r.raise_for_status()
    if len(r.content) < 3000:
        raise RuntimeError("Downloaded image is too small.")
    pil_from_bytes(r.content)
    return r.content


def challenge_detected(page):
    try:
        text = page.locator("body").inner_text(timeout=4000).lower()
    except Exception:
        return False
    return any(term in text for term in (
        "verification required", "slide right to secure your access", "verify you are human",
        "captcha", "security check", "access denied",
    ))


def wait_for_manual_verification(page, seconds=120):
    if not challenge_detected(page):
        return True
    st.warning("Shutterstock requested verification. Complete it in the visible Chrome window. The checker will continue automatically.")
    deadline = time.time() + seconds
    while time.time() < deadline:
        page.wait_for_timeout(1500)
        if not challenge_detected(page):
            page.wait_for_timeout(1500)
            return True
    return False


def dismiss_overlays(page):
    for selector in (
        'button:has-text("Accept all")', 'button:has-text("Accept All")',
        'button:has-text("Accept")', 'button:has-text("I agree")',
        '[id*="onetrust-accept" i]',
    ):
        try:
            loc = page.locator(selector).first
            if loc.count() and loc.is_visible():
                loc.click(timeout=1500)
                page.wait_for_timeout(500)
                return
        except Exception:
            pass


def safe_screenshot(page):
    try:
        return page.screenshot(full_page=False)
    except Exception:
        return None


def page_note(page):
    try:
        title = page.title()
    except Exception:
        title = ""
    return f"Page: {title or '(no title)'} | {getattr(page, 'url', '')}"


def set_file_inputs_everywhere(page, temp_path):
    for frame in page.frames:
        for selector in ('input[type="file"]', 'input[accept*="image" i]', 'input[name*="file" i]'):
            try:
                loc = frame.locator(selector)
                for i in range(min(loc.count(), 8)):
                    try:
                        loc.nth(i).set_input_files(temp_path, timeout=2500)
                        return True, f"file input in frame {frame.url or 'main'}"
                    except Exception:
                        pass
            except Exception:
                pass
    return False, ""


def click_text_search_by_image(page):
    script = """
    () => {
      const norm = s => (s || '').replace(/\s+/g, ' ').trim().toLowerCase();
      const all = [...document.querySelectorAll('button,[role="button"],a,label,div,span')];
      let best = null;
      for (const el of all) {
        const text = norm((el.innerText || '') + ' ' + (el.getAttribute('aria-label') || '') + ' ' + (el.getAttribute('title') || ''));
        if (text === 'search by image' || text.includes('search by image')) {
          const r = el.getBoundingClientRect();
          const visible = r.width > 0 && r.height > 0 && getComputedStyle(el).visibility !== 'hidden' && getComputedStyle(el).display !== 'none';
          if (!visible) continue;
          let target = el.closest('button,[role="button"],a,label') || el;
          best = target;
          break;
        }
      }
      if (!best) return {ok:false};
      best.click();
      return {ok:true, tag:best.tagName, text:norm(best.innerText || best.getAttribute('aria-label') || best.getAttribute('title') || '')};
    }
    """
    try:
        return page.evaluate(script)
    except Exception:
        return {"ok": False}


def click_near_search_camera(page):
    search = None
    candidates = page.locator('input[placeholder*="Search for images" i], input[placeholder*="Search" i], input[type="search"]')
    for i in range(min(candidates.count(), 12)):
        try:
            if candidates.nth(i).is_visible():
                search = candidates.nth(i)
                break
        except Exception:
            pass
    if search is None:
        return False, "no visible search input"

    try:
        box = search.bounding_box()
    except Exception:
        box = None
    if not box:
        return False, "no search box geometry"

    y = box["y"] + box["height"] / 2
    right = box["x"] + box["width"]

    for offset in (28, 42, 56, 70, 86, 104, 124, 146):
        x = right - offset
        try:
            result = page.evaluate(
                """({x,y}) => {
                  let el = document.elementFromPoint(x,y);
                  if (!el) return null;
                  let target = el.closest('button,[role="button"],a,label,[tabindex]') || el.parentElement || el;
                  const desc = (target.innerText || target.getAttribute('aria-label') || target.getAttribute('title') || target.tagName || '').trim();
                  target.click();
                  return desc;
                }""",
                {"x": x, "y": y},
            )
            page.wait_for_timeout(900)
            return True, f"clicked near search field offset {offset}px ({result})"
        except Exception:
            pass
    return False, "coordinate click failed"


def synthetic_drop(page, temp_path):
    try:
        raw = Path(temp_path).read_bytes()
        b64 = base64.b64encode(raw).decode("ascii")
        result = page.evaluate(
            """({b64,name}) => {
              const bytes = Uint8Array.from(atob(b64), c => c.charCodeAt(0));
              const file = new File([bytes], name, {type:'image/jpeg'});
              const dt = new DataTransfer();
              dt.items.add(file);
              const visible = el => {
                const r = el.getBoundingClientRect();
                return r.width > 0 && r.height > 0 && getComputedStyle(el).display !== 'none' && getComputedStyle(el).visibility !== 'hidden';
              };
              const norm = s => (s || '').replace(/\s+/g,' ').trim().toLowerCase();
              const els = [...document.querySelectorAll('[role="dialog"],input,button,label,div,section,main')].filter(visible);
              const targets = els.filter(el => {
                const t = norm(el.innerText || '') + ' ' + norm(el.getAttribute('aria-label') || '');
                return t.includes('drag and drop') || t.includes('drop the image') || t.includes('upload image') || t.includes('search by image');
              });
              const dispatchTo = targets.length ? targets.slice(0,8) : [document.body];
              for (const target of dispatchTo) {
                for (const type of ['dragenter','dragover','drop']) {
                  const ev = new DragEvent(type, {bubbles:true, cancelable:true, dataTransfer:dt});
                  target.dispatchEvent(ev);
                }
              }
              return {count:dispatchTo.length, text: norm(dispatchTo[0]?.innerText || '').slice(0,120)};
            }""",
            {"b64": b64, "name": "image.jpg"},
        )
        page.wait_for_timeout(1800)
        return True, f"synthetic drag/drop {result}"
    except Exception as e:
        return False, f"synthetic drop failed: {e}"


def trigger_search_by_image(page, temp_path):
    # A file input may already exist.
    ok, method = set_file_inputs_everywhere(page, temp_path)
    if ok:
        return True, method

    # Use Shutterstock's current search results page: its top search bar exposes Search by image.
    clicked = click_text_search_by_image(page)
    if clicked.get("ok"):
        page.wait_for_timeout(1300)
        ok, method = set_file_inputs_everywhere(page, temp_path)
        if ok:
            return True, f"text control -> {method}"
        ok, drop_method = synthetic_drop(page, temp_path)
        if ok:
            page.wait_for_timeout(1600)
            if page.url != "https://www.shutterstock.com/search":
                return True, f"text control -> {drop_method}"
            ok2, method2 = set_file_inputs_everywhere(page, temp_path)
            if ok2:
                return True, f"text control -> {method2}"

    # Try Playwright semantic locators.
    locators = [
        page.get_by_text(re.compile(r"search\s*by\s*image", re.I)),
        page.get_by_role("button", name=re.compile(r"search\s*by\s*image|camera|upload", re.I)),
        page.locator('[aria-label*="search by image" i], [title*="search by image" i], [aria-label*="camera" i], [title*="camera" i]'),
    ]
    for loc in locators:
        try:
            for i in range(min(loc.count(), 12)):
                el = loc.nth(i)
                if not el.is_visible():
                    continue
                try:
                    with page.expect_file_chooser(timeout=1800) as chooser_info:
                        el.click(force=True, timeout=2200)
                    chooser_info.value.set_files(temp_path)
                    return True, "semantic control / file chooser"
                except Exception:
                    try:
                        el.click(force=True, timeout=2200)
                    except Exception:
                        pass
                    page.wait_for_timeout(1000)
                    ok, method = set_file_inputs_everywhere(page, temp_path)
                    if ok:
                        return True, f"semantic control -> {method}"
                    ok, drop_method = synthetic_drop(page, temp_path)
                    if ok:
                        return True, f"semantic control -> {drop_method}"
        except Exception:
            pass

    # Click the actual element under likely camera coordinates and retry input/drop.
    clicked_ok, click_method = click_near_search_camera(page)
    if clicked_ok:
        page.wait_for_timeout(1200)
        ok, method = set_file_inputs_everywhere(page, temp_path)
        if ok:
            return True, f"{click_method} -> {method}"
        ok, drop_method = synthetic_drop(page, temp_path)
        if ok:
            return True, f"{click_method} -> {drop_method}"

    return False, "No working Search by Image control/file input/drop zone was detected"


def wait_for_results(page):
    try:
        page.wait_for_function(
            """() => [...document.querySelectorAll('a[href]')].some(a =>
              a.href.includes('/image-photo/') || a.href.includes('/image-vector/') ||
              a.href.includes('/image-illustration/') || a.href.includes('/image-3d/'))""",
            timeout=40000,
        )
    except PlaywrightTimeoutError:
        pass


def collect_candidates(page):
    page.wait_for_timeout(2200)
    for _ in range(4):
        page.mouse.wheel(0, 900)
        page.wait_for_timeout(450)

    selectors = [
        'a[href*="/image-photo/"] img', 'a[href*="/image-vector/"] img',
        'a[href*="/image-illustration/"] img', 'a[href*="/image-3d/"] img',
    ]
    results, seen = [], set()
    for selector in selectors:
        locs = page.locator(selector)
        for i in range(min(locs.count(), MAX_RESULTS)):
            img = locs.nth(i)
            try:
                href = clean_url(img.evaluate("el => { const a=el.closest('a'); return a ? a.href : ''; }"))
                if not href or href in seen:
                    continue
                img.scroll_into_view_if_needed(timeout=3000)
                shot = img.screenshot(timeout=5000)
                if len(shot) < 1000:
                    continue
                seen.add(href)
                results.append({"url": href, "bytes": shot})
                if len(results) >= MAX_RESULTS:
                    return results
            except Exception:
                continue
    return results


def check_shutterstock(page, upload_bytes):
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp:
            tmp.write(upload_bytes)
            temp_path = tmp.name

        page.goto("https://www.shutterstock.com/search", wait_until="domcontentloaded", timeout=90000)
        page.wait_for_timeout(4000)

        if not wait_for_manual_verification(page):
            return "MANUAL CHECK", "", 0.0, "Verification was not completed. " + page_note(page), safe_screenshot(page)

        dismiss_overlays(page)
        page.wait_for_timeout(700)

        uploaded, method = trigger_search_by_image(page, temp_path)
        if not uploaded:
            return "AUTOMATION ERROR", "", 0.0, method + ". " + page_note(page), safe_screenshot(page)

        wait_for_results(page)

        if not wait_for_manual_verification(page):
            return "MANUAL CHECK", "", 0.0, "Verification appeared after upload. " + page_note(page), safe_screenshot(page)

        candidates = collect_candidates(page)
        if not candidates:
            return "NOT FOUND", "", 0.0, f"Upload method: {method}. No Shutterstock result cards detected. " + page_note(page), safe_screenshot(page)

        best_url, best_score = "", 0.0
        for candidate in candidates:
            score = similarity(upload_bytes, candidate["bytes"])
            if score > best_score:
                best_url, best_score = candidate["url"], score
            if score >= 0.97:
                break

        notes = f"Upload method: {method}. Compared {len(candidates)} results. Best similarity: {best_score:.3f}."
        if best_score >= MATCH_THRESHOLD:
            return "MATCH FOUND", best_url, best_score, notes, None
        if best_score >= POSSIBLE_THRESHOLD:
            return "POSSIBLE MATCH", best_url, best_score, notes, None
        return "NOT FOUND", "", best_score, notes, None
    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass


def make_excel(results):
    out = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = ["Article URL", "Article Title", "Image Position", "Article Image URL", "Status", "Shutterstock Image URL", "Confidence", "Notes"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in results:
        ws.append([
            r.get("article_url", ""), r.get("article_title", ""), r.get("position", ""), r.get("image_url", ""),
            r.get("status", ""), r.get("shutterstock_url", ""), r.get("confidence", 0), r.get("notes", ""),
        ])
        row = ws.max_row
        for col in (1, 4, 6):
            value = ws.cell(row=row, column=col).value
            if value:
                ws.cell(row=row, column=col).hyperlink = value
                ws.cell(row=row, column=col).style = "Hyperlink"
    ws.freeze_panes = "A2"
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def urls_from_excel(uploaded):
    df = pd.read_excel(uploaded)
    names = {str(c).strip().lower(): c for c in df.columns}
    col = names.get("address") or names.get("url")
    if col is None:
        raise ValueError('Excel must contain a column named "Address" or "URL".')
    return [str(v).strip() for v in df[col].dropna().tolist() if str(v).strip().startswith("http")]


st.title("🔎 Shutterstock Article Image Checker — Local V2")
st.caption("Uses your visible Chrome browser. No Shutterstock API is required.")

mode = st.radio("Input method", ["Single article", "Paste URLs", "Upload Excel"], horizontal=True)
single_url = ""
multi_urls = ""
uploaded_excel = None
if mode == "Single article":
    single_url = st.text_input("Article URL", placeholder="https://www.bayut.com/mybayut/...")
elif mode == "Paste URLs":
    multi_urls = st.text_area("Paste one article URL per line", height=180)
else:
    uploaded_excel = st.file_uploader('Upload Excel with an "Address" or "URL" column', type=["xlsx"])

if st.button("Start Check", type="primary", use_container_width=True):
    if os.name != "nt":
        st.error("This local version is intended for Windows.")
        st.stop()

    if mode == "Single article":
        urls = [single_url.strip()] if single_url.strip() else []
    elif mode == "Paste URLs":
        urls = [x.strip() for x in multi_urls.splitlines() if x.strip().startswith("http")]
    else:
        urls = urls_from_excel(uploaded_excel) if uploaded_excel else []

    if not urls:
        st.warning("Add at least one article URL.")
        st.stop()

    results = []
    progress = st.progress(0, text="Starting...")
    status_box = st.empty()
    live = st.empty()

    with sync_playwright() as p:
        context = p.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            channel="chrome",
            headless=False,
            viewport=None,
            args=["--start-maximized"],
        )
        pages = context.pages
        article_page = pages[0] if pages else context.new_page()
        shutter_page = context.new_page()

        try:
            for article_index, article_url in enumerate(urls, start=1):
                status_box.info(f"Article {article_index}/{len(urls)} — extracting images...")
                try:
                    title, images = extract_article_images(article_page, article_url)
                except Exception as e:
                    results.append({"article_url": article_url, "article_title": "", "position": "", "image_url": "", "status": "ARTICLE ERROR", "shutterstock_url": "", "confidence": 0, "notes": str(e)})
                    continue

                for image_index, image in enumerate(images, start=1):
                    status_box.info(f"Article {article_index}/{len(urls)} — image {image_index}/{len(images)} — checking Shutterstock...")
                    try:
                        original = download_image(image["src"])
                        normalized = normalize_image(original)
                        status, shutter_url, score, notes, debug = check_shutterstock(shutter_page, normalized)
                        results.append({
                            "article_url": article_url, "article_title": title, "position": image["position"],
                            "image_url": image["src"], "status": status, "shutterstock_url": shutter_url,
                            "confidence": round(score * 100, 1), "notes": notes, "_preview": normalized, "_debug": debug,
                        })
                    except Exception as e:
                        results.append({"article_url": article_url, "article_title": title, "position": image["position"], "image_url": image["src"], "status": "IMAGE ERROR", "shutterstock_url": "", "confidence": 0, "notes": str(e)})

                    live_rows = [{k: v for k, v in r.items() if not k.startswith("_")} for r in results]
                    live.dataframe(pd.DataFrame(live_rows), use_container_width=True, hide_index=True)

                progress.progress(article_index / len(urls), text=f"Completed {article_index}/{len(urls)} articles")
        finally:
            context.close()

    status_box.success("Check completed.")
    display = [{k: v for k, v in r.items() if not k.startswith("_")} for r in results]
    df = pd.DataFrame(display)
    st.subheader("Results")
    st.dataframe(df, use_container_width=True, hide_index=True, column_config={
        "article_url": st.column_config.LinkColumn("Article URL"),
        "image_url": st.column_config.LinkColumn("Article Image URL"),
        "shutterstock_url": st.column_config.LinkColumn("Shutterstock Image URL"),
        "confidence": st.column_config.NumberColumn("Confidence", format="%.1f%%"),
    })

    st.subheader("Image Review")
    for i, row in enumerate(results, start=1):
        if "_preview" not in row:
            continue
        expanded = row["status"] in ("MATCH FOUND", "POSSIBLE MATCH", "AUTOMATION ERROR", "MANUAL CHECK", "NOT FOUND")
        with st.expander(f"{i}. {row['position']} — {row['status']}", expanded=expanded):
            c1, c2 = st.columns(2)
            with c1:
                st.image(row["_preview"], caption="Article image", use_container_width=True)
            with c2:
                st.write(f"**{row['status']}**")
                st.write(f"Confidence: **{row['confidence']:.1f}%**")
                st.write(row.get("notes", ""))
                if row.get("shutterstock_url"):
                    st.link_button("Open Shutterstock match", row["shutterstock_url"])
                if row.get("_debug"):
                    st.write("**Shutterstock page at failure:**")
                    st.image(row["_debug"], use_container_width=True)

    st.download_button(
        "Download Excel Report", data=make_excel(results), file_name="shutterstock_check_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True,
    )
