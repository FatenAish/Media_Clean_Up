import io
import os
import re
import time
import tempfile
from pathlib import Path
from urllib.parse import urljoin, urlsplit, urlunsplit

import imagehash
import pandas as pd
import requests
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font
from PIL import Image, ImageOps
from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError

st.set_page_config(page_title="Shutterstock Article Image Checker", page_icon="🔎", layout="wide")

BASE_DIR = Path(__file__).resolve().parent
PROFILE_DIR = BASE_DIR / "work" / "shutterstock_chrome_profile"
PROFILE_DIR.mkdir(parents=True, exist_ok=True)

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/153.0.0.0 Safari/537.36"
)

MAX_RESULTS = 40
MATCH_THRESHOLD = 0.86
POSSIBLE_THRESHOLD = 0.76


def clean_url(url):
    url = str(url or "").strip()
    if not url:
        return ""
    p = urlsplit(url)
    return urlunsplit((p.scheme, p.netloc, p.path, "", ""))


def pil_from_bytes(data):
    return Image.open(io.BytesIO(data)).convert("RGB")


def normalize_image(data):
    img = ImageOps.exif_transpose(pil_from_bytes(data))
    img.thumbnail((3900, 3900), Image.Resampling.LANCZOS)
    for quality in (92, 88, 84, 80, 76, 72, 68):
        out = io.BytesIO()
        img.save(out, "JPEG", quality=quality, optimize=True)
        value = out.getvalue()
        if len(value) < 4_800_000:
            return value
    img.thumbnail((3000, 3000), Image.Resampling.LANCZOS)
    out = io.BytesIO()
    img.save(out, "JPEG", quality=65, optimize=True)
    return out.getvalue()


def crop_variant(img, ratio=0.84):
    w, h = img.size
    nw, nh = int(w * ratio), int(h * ratio)
    left = max(0, (w - nw) // 2)
    top = max(0, (h - nh) // 2)
    return img.crop((left, top, left + nw, top + nh))


def similarity(a_bytes, b_bytes):
    try:
        a = ImageOps.exif_transpose(pil_from_bytes(a_bytes))
        b = ImageOps.exif_transpose(pil_from_bytes(b_bytes))
        ac = crop_variant(a)
        bc = crop_variant(b)
        pairs = [
            (imagehash.phash(a, hash_size=16), imagehash.phash(b, hash_size=16)),
            (imagehash.dhash(a, hash_size=16), imagehash.dhash(b, hash_size=16)),
            (imagehash.whash(a, hash_size=16), imagehash.whash(b, hash_size=16)),
            (imagehash.phash(ac, hash_size=16), imagehash.phash(bc, hash_size=16)),
            (imagehash.dhash(ac, hash_size=16), imagehash.dhash(bc, hash_size=16)),
        ]
        scores = [max(0.0, 1.0 - ((x - y) / 256.0)) for x, y in pairs]
        scores.sort(reverse=True)
        return min(1.0, 0.55 * scores[0] + 0.25 * scores[1] + 0.20 * (sum(scores) / len(scores)))
    except Exception:
        return 0.0


def extract_article_images(page, article_url):
    page.goto(article_url, wait_until="domcontentloaded", timeout=90000)
    page.wait_for_timeout(2500)
    title = page.title()

    for _ in range(8):
        page.mouse.wheel(0, 1000)
        page.wait_for_timeout(250)

    candidates = []
    og = page.locator('meta[property="og:image"]')
    if og.count():
        src = og.get_attribute("content")
        if src:
            candidates.append({"src": urljoin(article_url, src), "width": 1200, "height": 600, "position": "Feature"})

    selectors = ["article img", "main article img", ".entry-content img", ".post-content img", ".article-content img"]
    selector = next((s for s in selectors if page.locator(s).count()), "main img")
    raw = page.locator(selector).evaluate_all(
        """imgs => imgs.map((img, idx) => ({
            src: img.currentSrc || img.src || img.getAttribute('data-src') || img.getAttribute('data-lazy-src') || img.getAttribute('data-original') || '',
            width: img.naturalWidth || img.width || 0,
            height: img.naturalHeight || img.height || 0,
            idx: idx + 1
        }))"""
    )

    for item in raw:
        src = str(item.get("src") or "").strip()
        if src and not src.startswith("data:"):
            candidates.append({
                "src": urljoin(article_url, src),
                "width": int(item.get("width") or 0),
                "height": int(item.get("height") or 0),
                "position": f"Body {item.get('idx', '')}",
            })

    excluded = (
        "logo", "icon", "avatar", "emoji", "favicon", "sprite", "placeholder",
        "author", "profile", "social", "facebook", "instagram", "linkedin",
        "youtube", "twitter", "whatsapp",
    )

    seen = set()
    final = []
    for item in candidates:
        src = clean_url(item["src"])
        if not src or src in seen:
            continue
        if any(term in src.lower() for term in excluded):
            continue
        if item["width"] and item["width"] < 450:
            continue
        if item["height"] and item["height"] < 220:
            continue
        seen.add(src)
        item["src"] = src
        final.append(item)

    return title, final


def download_image(url):
    r = requests.get(
        url,
        headers={"User-Agent": USER_AGENT, "Referer": "https://www.bayut.com/"},
        timeout=45,
    )
    r.raise_for_status()
    if len(r.content) < 3000:
        raise RuntimeError("Downloaded image is too small.")
    pil_from_bytes(r.content)
    return r.content


def challenge_detected(page):
    try:
        text = page.locator("body").inner_text(timeout=4000).lower()
    except Exception:
        return False
    terms = (
        "verification required",
        "slide right to secure your access",
        "verify you are human",
        "captcha",
        "security check",
        "access denied",
    )
    return any(term in text for term in terms)


def wait_for_manual_verification(page, seconds=120):
    if not challenge_detected(page):
        return True
    st.warning(
        "Shutterstock requested verification. Complete the slider in the visible Chrome window. "
        "The checker will continue automatically after verification."
    )
    deadline = time.time() + seconds
    while time.time() < deadline:
        page.wait_for_timeout(1500)
        if not challenge_detected(page):
            page.wait_for_timeout(1500)
            return True
    return False


def dismiss_overlays(page):
    selectors = (
        'button:has-text("Accept all")',
        'button:has-text("Accept All")',
        'button:has-text("Accept")',
        'button:has-text("I agree")',
        'button:has-text("Agree")',
        '[id*="onetrust-accept" i]',
    )
    for selector in selectors:
        try:
            loc = page.locator(selector).first
            if loc.count() and loc.is_visible():
                loc.click(timeout=1500)
                page.wait_for_timeout(500)
                return
        except Exception:
            pass


def safe_screenshot(page):
    try:
        return page.screenshot(full_page=False)
    except Exception:
        return None


def debug_page_note(page):
    try:
        title = page.title()
    except Exception:
        title = ""
    try:
        url = page.url
    except Exception:
        url = ""
    return f"Page: {title or '(no title)'} | {url or '(no URL)'}"


def direct_file_input(page):
    selectors = [
        'input[type="file"]',
        'input[accept*="image" i]',
        'input[name*="file" i]',
    ]
    for selector in selectors:
        loc = page.locator(selector)
        if loc.count():
            return loc.first
    return None


def set_existing_file_input(page, temp_path):
    upload = direct_file_input(page)
    if upload is None:
        return False
    try:
        upload.set_input_files(temp_path, timeout=4000)
        return True
    except Exception:
        return False


def click_for_file(page, locator, temp_path, label):
    try:
        with page.expect_file_chooser(timeout=2200) as chooser_info:
            locator.click(timeout=2500, force=True)
        chooser_info.value.set_files(temp_path)
        return True, f"{label} / file chooser"
    except Exception:
        pass

    try:
        locator.click(timeout=2500, force=True)
    except Exception:
        return False, ""

    for _ in range(5):
        page.wait_for_timeout(350)
        if set_existing_file_input(page, temp_path):
            return True, f"{label} / revealed file input"

    return False, ""


def upload_from_dialog(page, temp_path):
    if set_existing_file_input(page, temp_path):
        return True, "dialog file input"

    patterns = (
        "upload image",
        "upload an image",
        "choose image",
        "choose file",
        "browse",
        "select image",
        "select file",
        "drag and drop",
        "search by image",
    )

    controls = page.locator("button, [role='button'], label, a, [tabindex], [aria-label], [title], [data-testid]")
    for i in range(min(controls.count(), 260)):
        el = controls.nth(i)
        try:
            if not el.is_visible():
                continue
            attrs = [
                el.inner_text(timeout=200),
                el.get_attribute("aria-label"),
                el.get_attribute("title"),
                el.get_attribute("data-testid"),
                el.get_attribute("name"),
            ]
            hay = " ".join(str(v) for v in attrs if v).strip().lower()
        except Exception:
            continue

        if not hay or not any(p in hay for p in patterns):
            continue

        ok, method = click_for_file(page, el, temp_path, f"dialog control [{hay[:80]}]")
        if ok:
            return True, method

    return False, ""


def upload_image(page, temp_path):
    # Strategy 1: Shutterstock may already render a hidden file input.
    if set_existing_file_input(page, temp_path):
        return True, "direct file input"

    # Strategy 2: semantic Search-by-Image controls.
    semantic_locators = [
        page.get_by_role("button", name=re.compile(r"search\s*by\s*image", re.I)),
        page.get_by_role("button", name=re.compile(r"camera|image search|upload", re.I)),
        page.get_by_text(re.compile(r"^\s*Search by image\s*$", re.I)),
        page.locator('[aria-label*="search by image" i]'),
        page.locator('[title*="search by image" i]'),
        page.locator('[aria-label*="camera" i]'),
        page.locator('[title*="camera" i]'),
        page.locator('[data-testid*="camera" i]'),
        page.locator('[data-testid*="image-search" i]'),
        page.locator('[data-testid*="search-by-image" i]'),
    ]

    for loc in semantic_locators:
        try:
            for i in range(min(loc.count(), 8)):
                el = loc.nth(i)
                if not el.is_visible():
                    continue
                ok, method = click_for_file(page, el, temp_path, "semantic Search by Image control")
                if ok:
                    return True, method
                ok, method = upload_from_dialog(page, temp_path)
                if ok:
                    return True, f"semantic control -> {method}"
        except Exception:
            pass

    # Strategy 3: inspect the search-bar wrapper. Shutterstock's camera control is
    # commonly an icon-only control immediately to the right of the search field.
    search_candidates = page.locator(
        'input[placeholder*="Search for images" i], '
        'input[placeholder*="Search" i], '
        'input[type="search"]'
    )

    search = None
    for i in range(min(search_candidates.count(), 12)):
        candidate = search_candidates.nth(i)
        try:
            if candidate.is_visible():
                search = candidate
                break
        except Exception:
            pass

    if search is not None:
        try:
            wrapper = search.locator(
                "xpath=ancestor::*[self::form or @role='search' or contains(translate(@class,'SEARCH','search'),'search')][1]"
            )
            if wrapper.count():
                controls = wrapper.locator("button, [role='button'], label, [tabindex], svg")
                for i in reversed(range(min(controls.count(), 30))):
                    el = controls.nth(i)
                    try:
                        if not el.is_visible():
                            continue
                    except Exception:
                        continue
                    ok, method = click_for_file(page, el, temp_path, "search-bar control")
                    if ok:
                        return True, method
                    ok, method = upload_from_dialog(page, temp_path)
                    if ok:
                        return True, f"search-bar control -> {method}"
        except Exception:
            pass

        # Strategy 4: geometry-based controls around the right side of the search box.
        try:
            search_box = search.bounding_box()
        except Exception:
            search_box = None

        if search_box:
            nearby = page.locator("button, [role='button'], label, [tabindex], svg")
            candidates = []
            for i in range(min(nearby.count(), 320)):
                el = nearby.nth(i)
                try:
                    if not el.is_visible():
                        continue
                    box = el.bounding_box()
                except Exception:
                    continue
                if not box:
                    continue

                cy = box["y"] + box["height"] / 2
                sy = search_box["y"] + search_box["height"] / 2
                right_edge = search_box["x"] + search_box["width"]

                same_row = abs(cy - sy) <= max(40, search_box["height"])
                near_right = right_edge - 180 <= box["x"] <= right_edge + 100
                smallish = box["width"] <= 140 and box["height"] <= 100

                if same_row and near_right and smallish:
                    distance = abs((box["x"] + box["width"] / 2) - (right_edge - 45))
                    candidates.append((distance, el))

            candidates.sort(key=lambda x: x[0])
            for _, el in candidates[:20]:
                ok, method = click_for_file(page, el, temp_path, "near-search camera candidate")
                if ok:
                    return True, method
                ok, method = upload_from_dialog(page, temp_path)
                if ok:
                    return True, f"near-search control -> {method}"

            # Strategy 5: click likely camera-icon coordinates inside the search field.
            y = search_box["y"] + search_box["height"] / 2
            for offset in (24, 36, 48, 60, 72, 88, 104, 122, 142):
                x = search_box["x"] + search_box["width"] - offset
                try:
                    with page.expect_file_chooser(timeout=1800) as chooser_info:
                        page.mouse.click(x, y)
                    chooser_info.value.set_files(temp_path)
                    return True, f"camera coordinate {offset}px"
                except Exception:
                    pass

                page.wait_for_timeout(450)
                if set_existing_file_input(page, temp_path):
                    return True, f"camera coordinate {offset}px -> file input"

                ok, method = upload_from_dialog(page, temp_path)
                if ok:
                    return True, f"camera coordinate {offset}px -> {method}"

    # Final pass in case a modal was opened by one of the attempts above.
    return upload_from_dialog(page, temp_path)


def wait_for_results(page):
    try:
        page.wait_for_function(
            """() => {
                const body = (document.body.innerText || '').toLowerCase();
                const links = [...document.querySelectorAll('a[href]')];
                const hasAsset = links.some(a =>
                    a.href.includes('/image-photo/') ||
                    a.href.includes('/image-vector/') ||
                    a.href.includes('/image-illustration/') ||
                    a.href.includes('/image-3d/')
                );
                return hasAsset && (body.includes('results') || body.includes('search by image'));
            }""",
            timeout=40000,
        )
    except PlaywrightTimeoutError:
        pass


def collect_candidates(page):
    page.wait_for_timeout(2500)
    for _ in range(4):
        page.mouse.wheel(0, 900)
        page.wait_for_timeout(450)

    selectors = [
        'a[href*="/image-photo/"] img',
        'a[href*="/image-vector/"] img',
        'a[href*="/image-illustration/"] img',
        'a[href*="/image-3d/"] img',
    ]

    results = []
    seen = set()
    for selector in selectors:
        locs = page.locator(selector)
        for i in range(min(locs.count(), MAX_RESULTS)):
            img = locs.nth(i)
            try:
                href = clean_url(img.evaluate("el => { const a = el.closest('a'); return a ? a.href : ''; }"))
                if not href or href in seen:
                    continue
                img.scroll_into_view_if_needed(timeout=3000)
                shot = img.screenshot(timeout=5000)
                if len(shot) < 1000:
                    continue
                seen.add(href)
                results.append({"url": href, "bytes": shot})
                if len(results) >= MAX_RESULTS:
                    return results
            except Exception:
                continue
    return results


def check_shutterstock(page, upload_bytes):
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as tmp:
            tmp.write(upload_bytes)
            temp_path = tmp.name

        page.goto("https://www.shutterstock.com/", wait_until="domcontentloaded", timeout=90000)
        page.wait_for_timeout(4500)

        if not wait_for_manual_verification(page):
            return (
                "MANUAL CHECK", "", 0.0,
                "Verification was not completed within 120 seconds. " + debug_page_note(page),
                safe_screenshot(page),
            )

        dismiss_overlays(page)
        page.wait_for_timeout(1000)

        uploaded, method = upload_image(page, temp_path)
        if not uploaded:
            return (
                "AUTOMATION ERROR", "", 0.0,
                "Could not trigger Shutterstock Search by Image. " + debug_page_note(page),
                safe_screenshot(page),
            )

        wait_for_results(page)

        if not wait_for_manual_verification(page):
            return (
                "MANUAL CHECK", "", 0.0,
                "Verification was requested after upload and was not completed. " + debug_page_note(page),
                safe_screenshot(page),
            )

        candidates = collect_candidates(page)
        if not candidates:
            return (
                "NOT FOUND", "", 0.0,
                f"Upload succeeded via {method}, but no Shutterstock result cards were detected. " + debug_page_note(page),
                safe_screenshot(page),
            )

        best_url = ""
        best_score = 0.0
        for candidate in candidates:
            score = similarity(upload_bytes, candidate["bytes"])
            if score > best_score:
                best_url = candidate["url"]
                best_score = score
            if score >= 0.97:
                break

        notes = f"Uploaded via {method}. Compared {len(candidates)} Shutterstock results. Best similarity: {best_score:.3f}."

        if best_score >= MATCH_THRESHOLD:
            return "MATCH FOUND", best_url, best_score, notes, None
        if best_score >= POSSIBLE_THRESHOLD:
            return "POSSIBLE MATCH", best_url, best_score, notes, None
        return "NOT FOUND", "", best_score, notes, None

    finally:
        if temp_path and os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception:
                pass


def make_excel(results):
    out = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    headers = [
        "Article URL", "Article Title", "Image Position", "Article Image URL",
        "Status", "Shutterstock Image URL", "Confidence", "Notes",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for r in results:
        ws.append([
            r.get("article_url", ""),
            r.get("article_title", ""),
            r.get("position", ""),
            r.get("image_url", ""),
            r.get("status", ""),
            r.get("shutterstock_url", ""),
            r.get("confidence", 0),
            r.get("notes", ""),
        ])
        row = ws.max_row
        for col in (1, 4, 6):
            value = ws.cell(row=row, column=col).value
            if value:
                ws.cell(row=row, column=col).hyperlink = value
                ws.cell(row=row, column=col).style = "Hyperlink"

    ws.freeze_panes = "A2"
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def urls_from_excel(uploaded):
    df = pd.read_excel(uploaded)
    names = {str(c).strip().lower(): c for c in df.columns}
    col = names.get("address") or names.get("url")
    if col is None:
        raise ValueError('Excel must contain a column named "Address" or "URL".')
    return [str(v).strip() for v in df[col].dropna().tolist() if str(v).strip().startswith("http")]


st.title("🔎 Shutterstock Article Image Checker — Local")
st.caption("Uses your visible Chrome browser. No Shutterstock API is required.")
st.success(
    "The checker now targets Shutterstock's Search by Image control using semantic, modal, search-bar and camera-position fallbacks."
)

mode = st.radio("Input method", ["Single article", "Paste URLs", "Upload Excel"], horizontal=True)
single_url = ""
multi_urls = ""
uploaded_excel = None

if mode == "Single article":
    single_url = st.text_input("Article URL", placeholder="https://www.bayut.com/mybayut/...")
elif mode == "Paste URLs":
    multi_urls = st.text_area("Paste one article URL per line", height=180)
else:
    uploaded_excel = st.file_uploader('Upload Excel with an "Address" or "URL" column', type=["xlsx"])

if st.button("Start Check", type="primary", use_container_width=True):
    if os.name != "nt":
        st.error("This local version is intended for Windows with Python 3.12.")
        st.stop()

    try:
        if mode == "Single article":
            urls = [single_url.strip()] if single_url.strip() else []
        elif mode == "Paste URLs":
            urls = [x.strip() for x in multi_urls.splitlines() if x.strip().startswith("http")]
        else:
            urls = urls_from_excel(uploaded_excel) if uploaded_excel else []
    except Exception as e:
        st.error(str(e))
        st.stop()

    if not urls:
        st.warning("Add at least one article URL.")
        st.stop()

    results = []
    progress = st.progress(0, text="Starting...")
    status_box = st.empty()
    live = st.empty()

    try:
        with sync_playwright() as p:
            context = p.chromium.launch_persistent_context(
                user_data_dir=str(PROFILE_DIR),
                channel="chrome",
                headless=False,
                viewport=None,
                args=["--start-maximized"],
            )

            pages = context.pages
            article_page = pages[0] if pages else context.new_page()
            shutter_page = context.new_page()

            try:
                for article_index, article_url in enumerate(urls, start=1):
                    status_box.info(f"Article {article_index}/{len(urls)} — extracting images...")

                    try:
                        title, images = extract_article_images(article_page, article_url)
                    except Exception as e:
                        results.append({
                            "article_url": article_url,
                            "article_title": "",
                            "position": "",
                            "image_url": "",
                            "status": "ARTICLE ERROR",
                            "shutterstock_url": "",
                            "confidence": 0,
                            "notes": str(e),
                        })
                        continue

                    if not images:
                        results.append({
                            "article_url": article_url,
                            "article_title": title,
                            "position": "",
                            "image_url": "",
                            "status": "NO IMAGES FOUND",
                            "shutterstock_url": "",
                            "confidence": 0,
                            "notes": "",
                        })
                        continue

                    for image_index, image in enumerate(images, start=1):
                        status_box.info(
                            f"Article {article_index}/{len(urls)} — image {image_index}/{len(images)} — checking Shutterstock..."
                        )

                        try:
                            original = download_image(image["src"])
                            normalized = normalize_image(original)
                            status, shutter_url, score, notes, debug_screenshot = check_shutterstock(
                                shutter_page, normalized
                            )
                            results.append({
                                "article_url": article_url,
                                "article_title": title,
                                "position": image["position"],
                                "image_url": image["src"],
                                "status": status,
                                "shutterstock_url": shutter_url,
                                "confidence": round(score * 100, 1),
                                "notes": notes,
                                "_preview": normalized,
                                "_debug": debug_screenshot,
                            })
                        except Exception as e:
                            results.append({
                                "article_url": article_url,
                                "article_title": title,
                                "position": image["position"],
                                "image_url": image["src"],
                                "status": "IMAGE ERROR",
                                "shutterstock_url": "",
                                "confidence": 0,
                                "notes": str(e),
                            })

                        live_rows = [
                            {k: v for k, v in r.items() if not k.startswith("_")}
                            for r in results
                        ]
                        live.dataframe(pd.DataFrame(live_rows), use_container_width=True, hide_index=True)

                    progress.progress(
                        article_index / len(urls),
                        text=f"Completed {article_index}/{len(urls)} articles",
                    )
            finally:
                context.close()

    except Exception as e:
        st.error(f"Browser automation could not start: {e}")
        st.stop()

    status_box.success("Check completed.")
    progress.progress(1.0, text="Completed")

    display = [
        {k: v for k, v in r.items() if not k.startswith("_")}
        for r in results
    ]
    df = pd.DataFrame(display)

    st.subheader("Results")
    st.dataframe(
        df,
        use_container_width=True,
        hide_index=True,
        column_config={
            "article_url": st.column_config.LinkColumn("Article URL"),
            "image_url": st.column_config.LinkColumn("Article Image URL"),
            "shutterstock_url": st.column_config.LinkColumn("Shutterstock Image URL"),
            "confidence": st.column_config.NumberColumn("Confidence", format="%.1f%%"),
        },
    )

    st.subheader("Image Review")
    for i, row in enumerate(results, start=1):
        if "_preview" not in row:
            continue

        expanded = row["status"] in (
            "MATCH FOUND", "POSSIBLE MATCH", "AUTOMATION ERROR", "MANUAL CHECK"
        )
        with st.expander(f"{i}. {row['position']} — {row['status']}", expanded=expanded):
            c1, c2 = st.columns(2)
            with c1:
                st.image(row["_preview"], caption="Article image", use_container_width=True)
            with c2:
                st.write(f"**{row['status']}**")
                st.write(f"Confidence: **{row['confidence']:.1f}%**")
                st.write(row.get("notes", ""))
                if row.get("shutterstock_url"):
                    st.link_button("Open Shutterstock match", row["shutterstock_url"])
                else:
                    st.write("No Shutterstock URL saved.")
                if row.get("_debug"):
                    st.write("**Shutterstock page when the check failed:**")
                    st.image(row["_debug"], use_container_width=True)

    st.download_button(
        "Download Excel Report",
        data=make_excel(results),
        file_name="shutterstock_check_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
