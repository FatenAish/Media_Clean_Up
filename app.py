import os

# On Windows, keep using the local browser version because it can use the
# user's visible Chrome session. On Streamlit Cloud/Linux, use Shutterstock's
# official Computer Vision API instead of browser automation. This avoids the
# Verification Required page that blocks datacenter browsers.
if os.name == "nt":
    from app_local import *  # noqa: F401,F403
else:
    import base64
    import io
    import re
    from urllib.parse import urljoin, urlsplit, urlunsplit

    import imagehash
    import pandas as pd
    import requests
    import streamlit as st
    from bs4 import BeautifulSoup
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from PIL import Image, ImageOps
    from requests.auth import HTTPBasicAuth

    st.set_page_config(
        page_title="Shutterstock Article Image Checker",
        page_icon="🔎",
        layout="wide",
    )

    USER_AGENT = (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/153.0.0.0 Safari/537.36"
    )

    API_BASE = "https://api.shutterstock.com/v2"
    MAX_RESULTS = 50
    MATCH_THRESHOLD = 0.86
    POSSIBLE_THRESHOLD = 0.76


    def clean_url(url):
        url = str(url or "").strip()
        if not url:
            return ""
        p = urlsplit(url)
        return urlunsplit((p.scheme, p.netloc, p.path, "", ""))


    def pil_from_bytes(data):
        return Image.open(io.BytesIO(data)).convert("RGB")


    def normalize_image(data):
        img = ImageOps.exif_transpose(pil_from_bytes(data))
        img.thumbnail((3900, 3900), Image.Resampling.LANCZOS)

        for quality in (92, 88, 84, 80, 76, 72, 68):
            out = io.BytesIO()
            img.save(out, "JPEG", quality=quality, optimize=True)
            value = out.getvalue()
            if len(value) < 4_800_000:
                return value

        img.thumbnail((3000, 3000), Image.Resampling.LANCZOS)
        out = io.BytesIO()
        img.save(out, "JPEG", quality=65, optimize=True)
        return out.getvalue()


    def crop_variant(img, ratio=0.84):
        w, h = img.size
        nw, nh = int(w * ratio), int(h * ratio)
        left = max(0, (w - nw) // 2)
        top = max(0, (h - nh) // 2)
        return img.crop((left, top, left + nw, top + nh))


    def similarity(a_bytes, b_bytes):
        """Near-duplicate score that tolerates resizing and modest cropping."""
        try:
            a = ImageOps.exif_transpose(pil_from_bytes(a_bytes))
            b = ImageOps.exif_transpose(pil_from_bytes(b_bytes))
            ac = crop_variant(a)
            bc = crop_variant(b)

            pairs = [
                (imagehash.phash(a, hash_size=16), imagehash.phash(b, hash_size=16)),
                (imagehash.dhash(a, hash_size=16), imagehash.dhash(b, hash_size=16)),
                (imagehash.whash(a, hash_size=16), imagehash.whash(b, hash_size=16)),
                (imagehash.phash(ac, hash_size=16), imagehash.phash(bc, hash_size=16)),
                (imagehash.dhash(ac, hash_size=16), imagehash.dhash(bc, hash_size=16)),
            ]

            scores = [max(0.0, 1.0 - ((x - y) / 256.0)) for x, y in pairs]
            scores.sort(reverse=True)
            return min(
                1.0,
                0.55 * scores[0]
                + 0.25 * scores[1]
                + 0.20 * (sum(scores) / len(scores)),
            )
        except Exception:
            return 0.0


    def largest_from_srcset(srcset):
        if not srcset:
            return ""
        best_url = ""
        best_width = -1
        for part in srcset.split(","):
            bits = part.strip().split()
            if not bits:
                continue
            url = bits[0]
            width = 0
            if len(bits) > 1 and bits[1].endswith("w"):
                try:
                    width = int(bits[1][:-1])
                except Exception:
                    width = 0
            if width >= best_width:
                best_width = width
                best_url = url
        return best_url


    def extract_article_images(article_url):
        r = requests.get(
            article_url,
            headers={"User-Agent": USER_AGENT, "Accept-Language": "en-US,en;q=0.9"},
            timeout=45,
        )
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "html.parser")

        title = soup.title.get_text(" ", strip=True) if soup.title else ""
        candidates = []

        og = soup.find("meta", attrs={"property": "og:image"})
        if og and og.get("content"):
            candidates.append({
                "src": urljoin(article_url, og.get("content")),
                "position": "Feature",
                "width": 1200,
                "height": 600,
            })

        container = (
            soup.find("article")
            or soup.find("main")
            or soup.find(class_=re.compile(r"entry-content|post-content|article-content", re.I))
            or soup
        )

        body_number = 0
        for img in container.find_all("img"):
            src = (
                largest_from_srcset(img.get("srcset"))
                or largest_from_srcset(img.get("data-srcset"))
                or img.get("data-src")
                or img.get("data-lazy-src")
                or img.get("data-original")
                or img.get("src")
                or ""
            )
            src = str(src).strip()
            if not src or src.startswith("data:"):
                continue

            body_number += 1
            try:
                width = int(str(img.get("width") or "0").replace("px", ""))
            except Exception:
                width = 0
            try:
                height = int(str(img.get("height") or "0").replace("px", ""))
            except Exception:
                height = 0

            candidates.append({
                "src": urljoin(article_url, src),
                "position": f"Body {body_number}",
                "width": width,
                "height": height,
            })

        excluded = (
            "logo", "icon", "avatar", "emoji", "favicon", "sprite",
            "placeholder", "author", "profile", "social", "facebook",
            "instagram", "linkedin", "youtube", "twitter", "whatsapp",
        )

        final = []
        seen = set()

        for item in candidates:
            src = clean_url(item["src"])
            if not src or src in seen:
                continue
            if any(term in src.lower() for term in excluded):
                continue
            if item.get("width") and item["width"] < 450:
                continue
            if item.get("height") and item["height"] < 220:
                continue

            seen.add(src)
            item["src"] = src
            final.append(item)

        return title, final


    def download_image(url):
        r = requests.get(
            url,
            headers={"User-Agent": USER_AGENT, "Referer": "https://www.bayut.com/"},
            timeout=45,
        )
        r.raise_for_status()
        if len(r.content) < 3000:
            raise RuntimeError("Downloaded image is too small.")
        pil_from_bytes(r.content)
        return r.content


    def slugify(text):
        text = re.sub(r"[^a-zA-Z0-9]+", "-", str(text or "").lower()).strip("-")
        return text[:180] or "stock-photo"


    def asset_page_url(item):
        asset_id = str(item.get("id", "")).strip()
        description = slugify(item.get("description", ""))
        return f"https://www.shutterstock.com/image-photo/{description}-{asset_id}"


    def auth_for_requests(token, key, secret):
        headers = {
            "User-Agent": "Media-Clean-Up/1.0",
            "Accept": "application/json",
        }
        auth = None

        if token:
            headers["Authorization"] = f"Bearer {token.strip()}"
        elif key and secret:
            auth = HTTPBasicAuth(key.strip(), secret.strip())
        else:
            raise RuntimeError("Add a Shutterstock API token or API key + secret.")

        return headers, auth


    def shutterstock_cv_search(image_bytes, token, key, secret):
        headers, auth = auth_for_requests(token, key, secret)

        upload_headers = dict(headers)
        upload_headers["Content-Type"] = "application/json"

        upload_response = requests.post(
            f"{API_BASE}/cv/images",
            headers=upload_headers,
            auth=auth,
            json={"base64_image": base64.b64encode(image_bytes).decode("ascii")},
            timeout=90,
        )

        if upload_response.status_code in (401, 403):
            raise RuntimeError(
                "Shutterstock rejected the API credentials or Computer Vision access is not enabled for this API application."
            )
        upload_response.raise_for_status()

        upload_id = upload_response.json().get("upload_id")
        if not upload_id:
            raise RuntimeError("Shutterstock API did not return an upload_id.")

        similar_response = requests.get(
            f"{API_BASE}/cv/similar/images",
            headers=headers,
            auth=auth,
            params={
                "asset_id": upload_id,
                "per_page": MAX_RESULTS,
                "page": 1,
                "view": "full",
                "safe": "true",
            },
            timeout=90,
        )

        if similar_response.status_code in (401, 403):
            raise RuntimeError(
                "Shutterstock Computer Vision search is not enabled for these API credentials."
            )
        similar_response.raise_for_status()

        data = similar_response.json()
        return data.get("data", []), data.get("total_count", 0)


    def preview_url(item):
        assets = item.get("assets") or {}
        for name in ("preview_1500", "preview_1000", "preview", "huge_thumb", "mosaic"):
            value = assets.get(name) or {}
            if value.get("url"):
                return value["url"]
        return ""


    def compare_api_results(original, items):
        best_item = None
        best_score = 0.0
        compared = 0

        for item in items:
            p_url = preview_url(item)
            if not p_url:
                continue

            try:
                r = requests.get(
                    p_url,
                    headers={"User-Agent": USER_AGENT},
                    timeout=30,
                )
                r.raise_for_status()
                score = similarity(original, r.content)
            except Exception:
                continue

            compared += 1
            if score > best_score:
                best_score = score
                best_item = item

            if score >= 0.97:
                break

        if best_item is None:
            return "NOT FOUND", "", 0.0, 0, None

        url = asset_page_url(best_item)

        if best_score >= MATCH_THRESHOLD:
            status = "MATCH FOUND"
        elif best_score >= POSSIBLE_THRESHOLD:
            status = "POSSIBLE MATCH"
        else:
            status = "NOT FOUND"
            url = ""

        return status, url, best_score, compared, best_item


    def make_excel(results):
        out = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.title = "Results"

        headers = [
            "Article URL",
            "Article Title",
            "Image Position",
            "Article Image URL",
            "Status",
            "Shutterstock Image URL",
            "Shutterstock Asset ID",
            "Confidence",
            "API Results",
            "Notes",
        ]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)

        for row in results:
            ws.append([
                row.get("article_url", ""),
                row.get("article_title", ""),
                row.get("position", ""),
                row.get("image_url", ""),
                row.get("status", ""),
                row.get("shutterstock_url", ""),
                row.get("asset_id", ""),
                row.get("confidence", 0),
                row.get("api_results", 0),
                row.get("notes", ""),
            ])

            r = ws.max_row
            for col in (1, 4, 6):
                value = ws.cell(r, col).value
                if value:
                    ws.cell(r, col).hyperlink = value
                    ws.cell(r, col).style = "Hyperlink"

        ws.freeze_panes = "A2"
        wb.save(out)
        out.seek(0)
        return out.getvalue()


    def urls_from_excel(uploaded):
        df = pd.read_excel(uploaded)
        names = {str(c).strip().lower(): c for c in df.columns}
        col = names.get("address") or names.get("url")
        if col is None:
            raise ValueError('Excel must contain a column named "Address" or "URL".')
        return [
            str(v).strip()
            for v in df[col].dropna().tolist()
            if str(v).strip().startswith("http")
        ]


    def secret_value(name):
        try:
            return str(st.secrets.get(name, "") or "")
        except Exception:
            return ""


    st.title("🔎 Shutterstock Article Image Checker")
    st.caption(
        "Cloud mode uses Shutterstock's official Computer Vision API, so it does not depend on the blocked Shutterstock website browser."
    )

    with st.sidebar:
        st.header("Shutterstock API")
        st.write("Use either an API token or API key + secret.")

        default_token = secret_value("SHUTTERSTOCK_API_TOKEN")
        default_key = secret_value("SHUTTERSTOCK_API_KEY")
        default_secret = secret_value("SHUTTERSTOCK_API_SECRET")

        api_token = st.text_input(
            "API token",
            value=default_token,
            type="password",
            help="Bearer token for a Shutterstock API application with Computer Vision access.",
        )
        st.caption("OR")
        api_key = st.text_input("API key", value=default_key, type="password")
        api_secret = st.text_input("API secret", value=default_secret, type="password")

        st.info(
            "The reverse-image endpoints must be enabled for your Shutterstock API application. Browser verification is not used in this mode."
        )

        mode = st.radio(
            "Input method",
            ["Single article", "Paste URLs", "Upload Excel"],
        )

    single_url = ""
    multi_urls = ""
    uploaded_excel = None

    if mode == "Single article":
        single_url = st.text_input(
            "Article URL",
            placeholder="https://www.bayut.com/mybayut/...",
        )
    elif mode == "Paste URLs":
        multi_urls = st.text_area("Paste one article URL per line", height=180)
    else:
        uploaded_excel = st.file_uploader(
            'Upload Excel with an "Address" or "URL" column',
            type=["xlsx"],
        )

    if st.button("Start Check", type="primary", use_container_width=True):
        if not api_token and not (api_key and api_secret):
            st.error("Add Shutterstock API credentials in the sidebar first.")
            st.stop()

        try:
            if mode == "Single article":
                urls = [single_url.strip()] if single_url.strip() else []
            elif mode == "Paste URLs":
                urls = [
                    x.strip()
                    for x in multi_urls.splitlines()
                    if x.strip().startswith("http")
                ]
            else:
                urls = urls_from_excel(uploaded_excel) if uploaded_excel else []
        except Exception as e:
            st.error(str(e))
            st.stop()

        if not urls:
            st.warning("Add at least one article URL.")
            st.stop()

        results = []
        progress = st.progress(0, text="Starting...")
        status_box = st.empty()
        live = st.empty()

        stop_for_api_error = False

        for article_index, article_url in enumerate(urls, start=1):
            status_box.info(
                f"Article {article_index}/{len(urls)} — extracting images..."
            )

            try:
                title, images = extract_article_images(article_url)
            except Exception as e:
                results.append({
                    "article_url": article_url,
                    "article_title": "",
                    "position": "",
                    "image_url": "",
                    "status": "ARTICLE ERROR",
                    "shutterstock_url": "",
                    "asset_id": "",
                    "confidence": 0,
                    "api_results": 0,
                    "notes": str(e),
                })
                continue

            if not images:
                results.append({
                    "article_url": article_url,
                    "article_title": title,
                    "position": "",
                    "image_url": "",
                    "status": "NO IMAGES FOUND",
                    "shutterstock_url": "",
                    "asset_id": "",
                    "confidence": 0,
                    "api_results": 0,
                    "notes": "",
                })
                continue

            for image_index, image in enumerate(images, start=1):
                status_box.info(
                    f"Article {article_index}/{len(urls)} — "
                    f"image {image_index}/{len(images)} — Shutterstock API search..."
                )

                try:
                    original = download_image(image["src"])
                    normalized = normalize_image(original)
                    api_items, total_count = shutterstock_cv_search(
                        normalized,
                        api_token,
                        api_key,
                        api_secret,
                    )

                    status, shutter_url, score, compared, best_item = compare_api_results(
                        normalized,
                        api_items,
                    )

                    results.append({
                        "article_url": article_url,
                        "article_title": title,
                        "position": image["position"],
                        "image_url": image["src"],
                        "status": status,
                        "shutterstock_url": shutter_url,
                        "asset_id": (best_item or {}).get("id", ""),
                        "confidence": round(score * 100, 1),
                        "api_results": total_count,
                        "notes": f"Compared {compared} API result previews.",
                        "_preview": normalized,
                        "_shutter_preview": preview_url(best_item or {}),
                    })

                except Exception as e:
                    message = str(e)
                    results.append({
                        "article_url": article_url,
                        "article_title": title,
                        "position": image["position"],
                        "image_url": image["src"],
                        "status": "API ERROR",
                        "shutterstock_url": "",
                        "asset_id": "",
                        "confidence": 0,
                        "api_results": 0,
                        "notes": message,
                        "_preview": normalized if "normalized" in locals() else None,
                    })

                    if "Computer Vision" in message or "credentials" in message:
                        stop_for_api_error = True

                display_rows = [
                    {k: v for k, v in r.items() if not k.startswith("_")}
                    for r in results
                ]
                live.dataframe(
                    pd.DataFrame(display_rows),
                    use_container_width=True,
                    hide_index=True,
                )

                if stop_for_api_error:
                    break

            progress.progress(
                article_index / len(urls),
                text=f"Completed {article_index}/{len(urls)} articles",
            )

            if stop_for_api_error:
                break

        if stop_for_api_error:
            status_box.error(
                "Stopped because Shutterstock API access failed. Check the credentials and make sure Computer Vision access is enabled for the API application."
            )
        else:
            status_box.success("Check completed.")
            progress.progress(1.0, text="Completed")

        display_rows = [
            {k: v for k, v in r.items() if not k.startswith("_")}
            for r in results
        ]
        df = pd.DataFrame(display_rows)

        st.subheader("Results")
        st.dataframe(
            df,
            use_container_width=True,
            hide_index=True,
            column_config={
                "article_url": st.column_config.LinkColumn("Article URL"),
                "image_url": st.column_config.LinkColumn("Article Image URL"),
                "shutterstock_url": st.column_config.LinkColumn("Shutterstock Image URL"),
                "confidence": st.column_config.NumberColumn("Confidence", format="%.1f%%"),
            },
        )

        st.subheader("Image Review")
        for i, row in enumerate(results, start=1):
            if not row.get("_preview"):
                continue

            with st.expander(
                f"{i}. {row['position']} — {row['status']}",
                expanded=row["status"] in ("MATCH FOUND", "POSSIBLE MATCH", "API ERROR"),
            ):
                c1, c2 = st.columns(2)

                with c1:
                    st.image(
                        row["_preview"],
                        caption="Article image",
                        use_container_width=True,
                    )

                with c2:
                    st.write(f"**{row['status']}**")
                    st.write(f"Confidence: **{row['confidence']:.1f}%**")
                    st.write(row.get("notes", ""))

                    if row.get("_shutter_preview"):
                        st.image(
                            row["_shutter_preview"],
                            caption="Best Shutterstock candidate",
                            use_container_width=True,
                        )

                    if row.get("shutterstock_url"):
                        st.link_button(
                            "Open Shutterstock match",
                            row["shutterstock_url"],
                        )

        st.download_button(
            "Download Excel Report",
            data=make_excel(results),
            file_name="shutterstock_check_results.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
